"""
Views for Analytics and Reporting.
"""
import csv
from io import StringIO
from rest_framework import generics, viewsets, status, permissions
from rest_framework.response import Response
from rest_framework.views import APIView
from django.contrib.auth import get_user_model
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Avg, Sum, Q
from datetime import timedelta
from .models import AttendanceReport, AttendanceTrend
from .serializers import (
    AttendanceReportSerializer,
    AttendanceReportCreateSerializer, AttendanceTrendSerializer
)
from .services import ReportService
from apps.users.permissions import IsAdmin, IsLecturer, IsStudent, IsLecturerOrAdmin
from apps.attendance.models import AttendanceRecord, AttendanceSession
from apps.courses.models import Course, Enrollment

User = get_user_model()


class AttendanceReportViewSet(viewsets.ModelViewSet):
    """ViewSet for Attendance Reports."""
    permission_classes = [permissions.IsAuthenticated, IsLecturerOrAdmin]

    def get_queryset(self):
        user = self.request.user
        queryset = AttendanceReport.objects.all()

        if user.role == 'lecturer':
            queryset = queryset.filter(
                Q(generated_by=user) | Q(course__lecturer=user)
            )

        return queryset

    def get_serializer_class(self):
        if self.action == 'create':
            return AttendanceReportCreateSerializer
        return AttendanceReportSerializer

    def perform_create(self, serializer):
        report = serializer.save(generated_by=self.request.user)
        self._generate_report(report)

    def _generate_report(self, report):
        """Generate report data."""
        records = AttendanceRecord.objects.filter(
            session__date__gte=report.period_start,
            session__date__lte=report.period_end
        )

        if report.course:
            records = records.filter(session__course=report.course)
        if report.department:
            records = records.filter(session__course__department=report.department)

        total = records.count()
        present = records.filter(status='present').count()
        late = records.filter(status='late').count()
        absent = records.filter(status='absent').count()

        report.summary = {
            'total_records': total,
            'present': present,
            'late': late,
            'absent': absent,
            'attendance_rate': round((present + late) / total * 100, 1) if total > 0 else 0
        }
        report.status = 'completed'
        report.generated_at = timezone.now()
        report.save()


class AttendanceTrendViewSet(viewsets.ReadOnlyModelViewSet):
    """ViewSet for Attendance Trends."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = AttendanceTrendSerializer

    def get_queryset(self):
        queryset = AttendanceTrend.objects.all()

        course_id = self.request.query_params.get('course')
        department = self.request.query_params.get('department')
        granularity = self.request.query_params.get('granularity')

        if course_id:
            queryset = queryset.filter(course_id=course_id)
        if department:
            queryset = queryset.filter(department=department)
        if granularity:
            queryset = queryset.filter(granularity=granularity)

        return queryset


class AtRiskStudentsView(APIView):
    """Get students at risk of falling below attendance threshold."""
    permission_classes = [permissions.IsAuthenticated, IsLecturerOrAdmin]

    def get(self, request):
        course_id = request.query_params.get('course')
        threshold = int(request.query_params.get('threshold', 75))

        course = None
        if course_id:
            try:
                course = Course.objects.get(id=course_id)
                if request.user.role == 'lecturer' and course.lecturer != request.user:
                    return Response(
                        {"error": "Access denied"},
                        status=status.HTTP_403_FORBIDDEN
                    )
            except Course.DoesNotExist:
                return Response(
                    {"error": "Course not found"},
                    status=status.HTTP_404_NOT_FOUND
                )

        students = ReportService.get_at_risk_students(course=course, threshold=threshold)
        return Response(students)


class DashboardAnalyticsView(APIView):
    """Get analytics data for dashboards."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user

        if user.role == 'student':
            return self._student_analytics(user)
        elif user.role == 'lecturer':
            return self._lecturer_analytics(user)
        elif user.role == 'admin':
            return self._admin_analytics()

        return Response({})

    def _student_analytics(self, student):
        """Get analytics for student dashboard."""
        enrollments = Enrollment.objects.filter(student=student, is_active=True)

        total_sessions = sum(e.total_sessions for e in enrollments)
        attended_sessions = sum(e.attended_sessions for e in enrollments)
        overall_rate = round((attended_sessions / total_sessions * 100), 1) if total_sessions > 0 else 0

        at_risk_courses = [e.course.code for e in enrollments if e.is_at_risk]

        thirty_days_ago = timezone.now() - timedelta(days=30)
        records = AttendanceRecord.objects.filter(
            student=student,
            session__date__gte=thirty_days_ago.date()
        ).order_by('session__date')

        trend_data = []
        current_date = thirty_days_ago.date()
        while current_date <= timezone.now().date():
            day_records = records.filter(session__date=current_date)
            if day_records.exists():
                present = day_records.filter(status__in=['present', 'late']).count()
                total = day_records.count()
                rate = round((present / total * 100), 1) if total > 0 else 0
                trend_data.append({
                    'date': current_date.isoformat(),
                    'rate': rate
                })
            current_date += timedelta(days=1)

        return Response({
            'overall_attendance': overall_rate,
            'total_courses': enrollments.count(),
            'at_risk_courses': at_risk_courses,
            'attended_sessions': attended_sessions,
            'total_sessions': total_sessions,
            'trend': trend_data
        })

    def _lecturer_analytics(self, lecturer):
        """Get analytics for lecturer dashboard."""
        courses = Course.objects.filter(lecturer=lecturer)

        course_stats = []
        for course in courses:
            enrollments = course.enrollments.filter(is_active=True)
            sessions = course.attendance_sessions.all()

            avg_attendance = 0
            if sessions.exists():
                total_present = sum(s.present_count + s.late_count for s in sessions)
                total_expected = sum(s.total_enrolled for s in sessions)
                avg_attendance = round((total_present / total_expected * 100), 1) if total_expected > 0 else 0

            at_risk = enrollments.filter(
                attendance_percentage__lt=course.attendance_threshold
            ).count()

            course_stats.append({
                'course_id': str(course.id),
                'course_code': course.code,
                'course_name': course.name,
                'total_students': enrollments.count(),
                'average_attendance': avg_attendance,
                'at_risk_count': at_risk,
                'total_sessions': sessions.count()
            })

        active_sessions = AttendanceSession.objects.filter(
            lecturer=lecturer,
            is_active=True
        ).count()

        today_sessions = AttendanceSession.objects.filter(
            lecturer=lecturer,
            date=timezone.now().date()
        ).count()

        return Response({
            'courses': course_stats,
            'total_courses': courses.count(),
            'active_sessions': active_sessions,
            'today_sessions': today_sessions
        })

    def _admin_analytics(self):
        """Get analytics for admin dashboard."""
        total_students = User.objects.filter(role='student', is_active=True).count()
        total_lecturers = User.objects.filter(role='lecturer', is_active=True).count()
        total_courses = Course.objects.filter(is_active=True).count()

        # Average attendance: mean of enrollment percentages (uses total_lectures as denominator)
        active_enrollments = Enrollment.objects.filter(is_active=True)
        pcts = list(active_enrollments.values_list('attendance_percentage', flat=True))
        avg_attendance = round(sum(float(p) for p in pcts) / len(pcts), 1) if pcts else 0

        active_sessions = AttendanceSession.objects.filter(is_active=True).count()

        # At-risk: students whose OVERALL average across all enrollments is below 75%
        at_risk_count = 0
        for student in User.objects.filter(role='student', is_active=True):
            enrollments = active_enrollments.filter(student=student)
            if not enrollments.exists():
                continue
            avg = sum(float(p) for p in enrollments.values_list('attendance_percentage', flat=True)) / enrollments.count()
            if avg < 75:
                at_risk_count += 1

        # Attendance by programme (more meaningful than raw department field)
        from apps.courses.models import Programme
        attendance_by_department = []
        for prog in Programme.objects.all():
            prog_courses = Course.objects.filter(programme=prog, is_active=True)
            if not prog_courses.exists():
                continue
            prog_enrollments = active_enrollments.filter(course__in=prog_courses)
            student_count = prog_enrollments.values('student').distinct().count()
            if student_count == 0:
                continue
            prog_pcts = list(prog_enrollments.values_list('attendance_percentage', flat=True))
            prog_avg = round(sum(float(p) for p in prog_pcts) / len(prog_pcts), 1) if prog_pcts else 0
            attendance_by_department.append({
                'department': prog.name,
                'percentage': prog_avg,
                'student_count': student_count,
            })
        attendance_by_department.sort(key=lambda x: x['percentage'], reverse=True)

        # Weekly trends: actual attendance rate per week from records
        all_records = AttendanceRecord.objects.all()
        recent_trends = []
        for i in range(5, -1, -1):
            week_start = (timezone.now() - timedelta(weeks=i + 1)).date()
            week_end = (timezone.now() - timedelta(weeks=i)).date()
            week_records = all_records.filter(
                session__date__gte=week_start,
                session__date__lt=week_end,
            )
            total = week_records.count()
            present = week_records.filter(status__in=['present', 'late']).count()
            percentage = round((present / total * 100), 1) if total > 0 else 0
            recent_trends.append({
                'period': f'Week {6 - i}',
                'percentage': percentage,
                'sessionsAttended': present,
                'totalSessions': total,
            })

        return Response({
            'totalStudents': total_students,
            'totalLecturers': total_lecturers,
            'totalCourses': total_courses,
            'activeSessions': active_sessions,
            'averageAttendance': avg_attendance,
            'studentsAtRisk': at_risk_count,
            'attendanceByDepartment': attendance_by_department,
            'recentTrends': recent_trends,
        })


class CourseAnalyticsView(APIView):
    """Get detailed analytics for a specific course."""
    permission_classes = [permissions.IsAuthenticated, IsLecturerOrAdmin]

    def get(self, request, course_id):
        try:
            course = Course.objects.get(id=course_id)
        except Course.DoesNotExist:
            return Response(
                {"error": "Course not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        if request.user.role == 'lecturer' and course.lecturer != request.user:
            return Response(
                {"error": "Access denied"},
                status=status.HTTP_403_FORBIDDEN
            )

        sessions = course.attendance_sessions.all()
        enrollments = course.enrollments.filter(is_active=True)

        session_data = []
        for session in sessions.order_by('date'):
            session_data.append({
                'date': session.date.isoformat(),
                'attendance_rate': session.attendance_rate,
                'present': session.present_count,
                'late': session.late_count,
                'absent': session.absent_count
            })

        distribution = {
            'excellent': enrollments.filter(attendance_percentage__gte=90).count(),
            'good': enrollments.filter(attendance_percentage__gte=75, attendance_percentage__lt=90).count(),
            'at_risk': enrollments.filter(attendance_percentage__gte=50, attendance_percentage__lt=75).count(),
            'critical': enrollments.filter(attendance_percentage__lt=50).count()
        }

        top_students = enrollments.order_by('-attendance_percentage')[:5]
        bottom_students = enrollments.order_by('attendance_percentage')[:5]

        return Response({
            'course': {
                'id': str(course.id),
                'code': course.code,
                'name': course.name,
                'threshold': course.attendance_threshold
            },
            'summary': {
                'total_sessions': sessions.count(),
                'total_students': enrollments.count(),
                'average_attendance': round(
                    sum(float(e.attendance_percentage) for e in enrollments) / enrollments.count(), 1
                ) if enrollments.exists() else 0
            },
            'session_data': session_data,
            'distribution': distribution,
            'top_students': [
                {
                    'name': e.student.full_name,
                    'student_id': e.student.student_id,
                    'attendance': float(e.attendance_percentage)
                }
                for e in top_students
            ],
            'bottom_students': [
                {
                    'name': e.student.full_name,
                    'student_id': e.student.student_id,
                    'attendance': float(e.attendance_percentage)
                }
                for e in bottom_students
            ]
        })


class AttendanceExportView(APIView):
    """Export a comprehensive attendance report as CSV."""
    permission_classes = [permissions.IsAuthenticated, IsAdmin]

    def get(self, request):
        from apps.attendance.models import AttendanceSession
        User = get_user_model()
        now = timezone.now()

        output = StringIO()
        writer = csv.writer(output)

        # ── Summary Section ──────────────────────────────────────────────────
        writer.writerow(["BIOATTEND — ATTENDANCE REPORT"])
        writer.writerow([f"Generated: {now.strftime('%d %B %Y %H:%M')}"])
        writer.writerow([])

        total_students = User.objects.filter(role='student', is_active=True).count()
        total_lecturers = User.objects.filter(role='lecturer', is_active=True).count()
        total_courses = Course.objects.filter(is_active=True).count()
        active_enrollments = Enrollment.objects.filter(is_active=True)
        pcts = [float(p) for p in active_enrollments.values_list('attendance_percentage', flat=True)]
        avg_att = round(sum(pcts) / len(pcts), 1) if pcts else 0

        at_risk_count = 0
        for s in User.objects.filter(role='student', is_active=True):
            enrs = active_enrollments.filter(student=s)
            if not enrs.exists():
                continue
            avg = sum(float(p) for p in enrs.values_list('attendance_percentage', flat=True)) / enrs.count()
            if avg < 75:
                at_risk_count += 1

        writer.writerow(["INSTITUTION SUMMARY"])
        writer.writerow(["Metric", "Value"])
        writer.writerow(["Total Students", total_students])
        writer.writerow(["Total Lecturers", total_lecturers])
        writer.writerow(["Total Active Courses", total_courses])
        writer.writerow(["Overall Average Attendance", f"{avg_att}%"])
        writer.writerow(["Students At Risk (avg < 75%)", at_risk_count])
        writer.writerow([])

        # ── Course Summary ───────────────────────────────────────────────────
        writer.writerow(["COURSE SUMMARY"])
        writer.writerow([
            "Course Code", "Course Name", "Programme", "Year", "Semester",
            "Lecturer", "Total Students", "Sessions Held", "Avg Attendance %",
            "At Risk Count",
        ])
        for course in Course.objects.filter(is_active=True).select_related('programme', 'lecturer').order_by('code'):
            enrs = Enrollment.objects.filter(course=course, is_active=True)
            total_enrolled = enrs.count()
            sessions_held = AttendanceSession.objects.filter(course=course).count()
            course_pcts = [float(p) for p in enrs.values_list('attendance_percentage', flat=True)]
            course_avg = round(sum(course_pcts) / len(course_pcts), 1) if course_pcts else 0
            at_risk = enrs.filter(attendance_percentage__lt=course.attendance_threshold).count()
            lecturer_name = course.lecturer.get_full_name() if course.lecturer else "—"
            prog_name = course.programme.name if course.programme else "—"
            writer.writerow([
                course.code, course.name, prog_name,
                course.year_level or "—", course.semester_number or "—",
                lecturer_name, total_enrolled, sessions_held,
                f"{course_avg}%", at_risk,
            ])
        writer.writerow([])

        # ── Student Detail ───────────────────────────────────────────────────
        writer.writerow(["STUDENT ATTENDANCE DETAIL"])
        writer.writerow([
            "Student Name", "Student ID", "Email", "Cohort",
            "Course Code", "Course Name",
            "Sessions Held", "Sessions Attended", "Sessions Absent",
            "Attendance %", "Total Lectures", "Threshold",
            "Status",
        ])

        students = User.objects.filter(role='student', is_active=True).order_by('last_name', 'first_name')
        for student in students:
            cohort_label = str(student.cohort) if student.cohort else "—"
            enrs = (
                Enrollment.objects
                .filter(student=student, is_active=True)
                .select_related('course')
                .order_by('course__code')
            )
            if not enrs.exists():
                writer.writerow([
                    student.get_full_name(), student.student_id or "—",
                    student.email, cohort_label,
                    "—", "—", 0, 0, 0, "0.0%", 0, "75%", "No Enrollments",
                ])
                continue

            for enr in enrs:
                course = enr.course
                pct = float(enr.attendance_percentage)
                status_label = "At Risk" if pct < course.attendance_threshold else "On Track"
                absent = enr.total_sessions - enr.attended_sessions
                writer.writerow([
                    student.get_full_name(),
                    student.student_id or "—",
                    student.email,
                    cohort_label,
                    course.code,
                    course.name,
                    enr.total_sessions,
                    enr.attended_sessions,
                    absent,
                    f"{pct:.1f}%",
                    course.total_lectures,
                    f"{course.attendance_threshold}%",
                    status_label,
                ])

        # ── Response ─────────────────────────────────────────────────────────
        filename = f"attendance_report_{now.strftime('%Y%m%d_%H%M')}.csv"
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class CoursePDFReportView(APIView):
    """Generate a detailed PDF attendance report for a single course."""
    permission_classes = [permissions.IsAuthenticated, IsAdmin]

    def get(self, request, course_id):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        from io import BytesIO

        try:
            course = Course.objects.select_related('programme', 'lecturer', 'faculty').get(pk=course_id)
        except Course.DoesNotExist:
            return Response({'error': 'Course not found'}, status=404)

        now = timezone.now()
        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm,
            title=f"{course.code} Attendance Report",
        )

        styles = getSampleStyleSheet()
        primary = colors.HexColor('#4F46E5')
        danger  = colors.HexColor('#DC2626')
        success = colors.HexColor('#16A34A')
        light   = colors.HexColor('#F3F4F6')

        title_style   = ParagraphStyle('Title',   parent=styles['Title'],   fontSize=18, textColor=primary, spaceAfter=4)
        heading_style = ParagraphStyle('Heading', parent=styles['Heading2'], fontSize=12, textColor=primary, spaceBefore=14, spaceAfter=4)
        sub_style     = ParagraphStyle('Sub',     parent=styles['Normal'],  fontSize=9,  textColor=colors.HexColor('#6B7280'))
        normal_style  = styles['Normal']

        story = []

        # ── Header ────────────────────────────────────────────────────────────
        story.append(Paragraph("BioAttend — Course Attendance Report", title_style))
        story.append(Paragraph(f"Generated: {now.strftime('%d %B %Y, %H:%M')}", sub_style))
        story.append(HRFlowable(width="100%", thickness=2, color=primary, spaceAfter=10))

        # ── Course Info ───────────────────────────────────────────────────────
        story.append(Paragraph("Course Information", heading_style))
        info_data = [
            ["Course Code",  course.code,       "Programme",   course.programme.name if course.programme else "—"],
            ["Course Name",  course.name,        "Department",  course.faculty.name if course.faculty else (course.department or "—")],
            ["Lecturer",     course.lecturer.get_full_name() if course.lecturer else "—",
             "Year / Semester", f"Year {course.year_level or '—'} · Sem {course.semester_number or '—'}"],
            ["Threshold",    f"{course.attendance_threshold}%",
             "Total Lectures", str(course.total_lectures)],
        ]
        info_table = Table(info_data, colWidths=[3.5*cm, 5.5*cm, 3.5*cm, 5.5*cm])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (0,-1), light),
            ('BACKGROUND', (2,0), (2,-1), light),
            ('FONTNAME',   (0,0), (-1,-1), 'Helvetica'),
            ('FONTNAME',   (0,0), (0,-1),  'Helvetica-Bold'),
            ('FONTNAME',   (2,0), (2,-1),  'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,-1), 9),
            ('GRID',       (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
            ('PADDING',    (0,0), (-1,-1), 6),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.3*cm))

        # ── Session Summary ───────────────────────────────────────────────────
        sessions = AttendanceSession.objects.filter(course=course).order_by('date')
        enrollments = Enrollment.objects.filter(course=course, is_active=True).select_related('student')
        total_enrolled = enrollments.count()
        sessions_held  = sessions.count()
        pcts = [float(e.attendance_percentage) for e in enrollments]
        avg_pct = round(sum(pcts)/len(pcts), 1) if pcts else 0
        above = sum(1 for p in pcts if p >= course.attendance_threshold)
        below = total_enrolled - above

        story.append(Paragraph("Summary", heading_style))
        summary_data = [
            ["Enrolled Students", str(total_enrolled), "Sessions Held",   str(sessions_held)],
            ["Avg Attendance",    f"{avg_pct}%",        "Above Threshold", str(above)],
            ["Below Threshold",   str(below),           "",                ""],
        ]
        s_table = Table(summary_data, colWidths=[4*cm, 4*cm, 4*cm, 4*cm])
        s_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (0,-1), light),
            ('BACKGROUND', (2,0), (2,-1), light),
            ('FONTNAME',   (0,0), (-1,-1), 'Helvetica'),
            ('FONTNAME',   (0,0), (0,-1),  'Helvetica-Bold'),
            ('FONTNAME',   (2,0), (2,-1),  'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,-1), 9),
            ('GRID',       (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
            ('PADDING',    (0,0), (-1,-1), 6),
        ]))
        story.append(s_table)
        story.append(Spacer(1, 0.3*cm))

        # ── Student Attendance Table ───────────────────────────────────────────
        story.append(Paragraph("Student Attendance Records", heading_style))
        headers = ["#", "Student Name", "Student ID", "Sessions\nAttended", "Sessions\nHeld", "Attendance\n%", "Status"]
        rows = [headers]
        for i, enr in enumerate(enrollments.order_by('student__last_name', 'student__first_name'), 1):
            pct = float(enr.attendance_percentage)
            status_label = "At Risk" if pct < course.attendance_threshold else "On Track"
            rows.append([
                str(i),
                enr.student.get_full_name(),
                enr.student.student_id or "—",
                str(enr.attended_sessions),
                str(enr.total_sessions),
                f"{pct:.1f}%",
                status_label,
            ])

        col_widths = [1*cm, 5.5*cm, 3.5*cm, 2.2*cm, 2.2*cm, 2.2*cm, 2.2*cm]
        student_table = Table(rows, colWidths=col_widths, repeatRows=1)
        style_cmds = [
            ('BACKGROUND',  (0,0), (-1,0),  primary),
            ('TEXTCOLOR',   (0,0), (-1,0),  colors.white),
            ('FONTNAME',    (0,0), (-1,0),  'Helvetica-Bold'),
            ('FONTNAME',    (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE',    (0,0), (-1,-1), 8),
            ('ALIGN',       (3,0), (-1,-1), 'CENTER'),
            ('GRID',        (0,0), (-1,-1), 0.4, colors.HexColor('#E5E7EB')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, light]),
            ('PADDING',     (0,0), (-1,-1), 5),
            ('VALIGN',      (0,0), (-1,-1), 'MIDDLE'),
        ]
        # Colour At Risk cells red, On Track green
        for row_idx, enr in enumerate(enrollments.order_by('student__last_name', 'student__first_name'), 1):
            pct = float(enr.attendance_percentage)
            col = danger if pct < course.attendance_threshold else success
            style_cmds.append(('TEXTCOLOR', (6, row_idx), (6, row_idx), col))
            style_cmds.append(('FONTNAME',  (6, row_idx), (6, row_idx), 'Helvetica-Bold'))
        student_table.setStyle(TableStyle(style_cmds))
        story.append(student_table)
        story.append(Spacer(1, 0.4*cm))

        # ── Session Log ───────────────────────────────────────────────────────
        if sessions.exists():
            story.append(Paragraph("Session Log", heading_style))
            sess_headers = ["Date", "Title", "Type", "Room", "Start", "End", "Present", "Absent"]
            sess_rows = [sess_headers]
            for s in sessions:
                present = AttendanceRecord.objects.filter(session=s, status='present').count()
                absent  = total_enrolled - present
                sess_rows.append([
                    s.date.strftime('%d %b %Y') if s.date else "—",
                    (s.title or s.course.name)[:30],
                    s.session_type or "lecture",
                    s.room or "—",
                    s.start_time.strftime('%H:%M') if s.start_time else "—",
                    s.end_time.strftime('%H:%M')   if s.end_time   else "—",
                    str(present),
                    str(absent),
                ])
            sess_table = Table(sess_rows, colWidths=[2.5*cm, 4.5*cm, 1.8*cm, 2*cm, 1.5*cm, 1.5*cm, 1.6*cm, 1.6*cm], repeatRows=1)
            sess_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0),  primary),
                ('TEXTCOLOR',  (0,0), (-1,0),  colors.white),
                ('FONTNAME',   (0,0), (-1,0),  'Helvetica-Bold'),
                ('FONTNAME',   (0,1), (-1,-1), 'Helvetica'),
                ('FONTSIZE',   (0,0), (-1,-1), 8),
                ('ALIGN',      (6,0), (7,-1),  'CENTER'),
                ('GRID',       (0,0), (-1,-1), 0.4, colors.HexColor('#E5E7EB')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, light]),
                ('PADDING',    (0,0), (-1,-1), 5),
            ]))
            story.append(sess_table)

        doc.build(story)
        buf.seek(0)
        fname = f"{course.code}_attendance_{now.strftime('%Y%m%d')}.pdf"
        resp = HttpResponse(buf.read(), content_type='application/pdf')
        resp['Content-Disposition'] = f'attachment; filename="{fname}"'
        return resp


class ReportFiltersView(APIView):
    """Return distinct filter options for the threshold Excel report."""
    permission_classes = [permissions.IsAuthenticated, IsAdmin]

    def get(self, request):
        from apps.courses.models import AcademicYear, Faculty, Programme

        academic_years = list(
            AcademicYear.objects.order_by('-start_year').values('id', 'label')
        )
        faculties = list(
            Faculty.objects.order_by('name').values('id', 'name')
        )
        year_levels = (
            Course.objects.filter(is_active=True, year_level__isnull=False)
            .order_by('year_level')
            .values_list('year_level', flat=True)
            .distinct()
        )
        semester_numbers = (
            Course.objects.filter(is_active=True, semester_number__isnull=False)
            .order_by('semester_number')
            .values_list('semester_number', flat=True)
            .distinct()
        )
        courses = list(
            Course.objects.filter(is_active=True)
            .order_by('code')
            .values('id', 'code', 'name')
        )
        return Response({
            'academic_years': academic_years,
            'faculties': faculties,
            'year_levels': list(year_levels),
            'semester_numbers': list(semester_numbers),
            'courses': courses,
        })


class ThresholdExcelExportView(APIView):
    """Export per-course student lists (above & below threshold) as Excel."""
    permission_classes = [permissions.IsAuthenticated, IsAdmin]

    def get(self, request):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from io import BytesIO

        now = timezone.now()
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        # ── Query filters from request params ──────────────────────────────
        academic_year_id = request.query_params.get('academic_year')   # AcademicYear id
        faculty_id       = request.query_params.get('faculty')          # Faculty id
        course_id        = request.query_params.get('course')           # Course id
        year_level       = request.query_params.get('year_level')
        semester_number  = request.query_params.get('semester_number')

        courses_qs = Course.objects.filter(is_active=True).select_related('programme', 'lecturer', 'faculty').order_by('code')

        if course_id:
            courses_qs = courses_qs.filter(id=course_id)
        else:
            if academic_year_id:
                # Filter courses whose enrolled students belong to cohorts with this intake year
                from apps.courses.models import Cohort
                cohort_ids = Cohort.objects.filter(intake_year_id=academic_year_id).values_list('id', flat=True)
                courses_qs = courses_qs.filter(
                    enrollments__student__cohort_id__in=cohort_ids,
                    enrollments__is_active=True
                ).distinct()
            if faculty_id:
                courses_qs = courses_qs.filter(faculty_id=faculty_id)
            if year_level:
                courses_qs = courses_qs.filter(year_level=year_level)
            if semester_number:
                courses_qs = courses_qs.filter(semester_number=semester_number)

        # Styles
        hdr_font   = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        above_fill = PatternFill('solid', fgColor='16A34A')  # green header
        below_fill = PatternFill('solid', fgColor='DC2626')  # red header
        meta_fill  = PatternFill('solid', fgColor='4F46E5')  # indigo header
        even_fill  = PatternFill('solid', fgColor='F3F4F6')
        center     = Alignment(horizontal='center', vertical='center')
        left       = Alignment(horizontal='left',   vertical='center')
        thin       = Side(style='thin', color='D1D5DB')
        border     = Border(left=thin, right=thin, top=thin, bottom=thin)

        def _hdr_row(ws, row, values, fill):
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font      = hdr_font
                cell.fill      = fill
                cell.alignment = center
                cell.border    = border

        def _data_row(ws, row, values, is_even=False):
            bg = PatternFill('solid', fgColor='FFFFFF') if not is_even else even_fill
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font      = Font(name='Calibri', size=10)
                cell.fill      = bg
                cell.alignment = left if col in (2, 3) else center
                cell.border    = border

        for course in courses_qs:
            enrollments = (
                Enrollment.objects
                .filter(course=course, is_active=True)
                .select_related('student')
                .order_by('student__last_name', 'student__first_name')
            )
            if not enrollments.exists():
                continue

            threshold = course.attendance_threshold
            above = [e for e in enrollments if float(e.attendance_percentage) >= threshold]
            below = [e for e in enrollments if float(e.attendance_percentage) <  threshold]

            # Sheet name max 31 chars
            sheet_name = course.code[:31]
            ws = wb.create_sheet(title=sheet_name)
            ws.sheet_view.showGridLines = False

            # ── Meta header ──
            ws.merge_cells('A1:G1')
            title_cell = ws['A1']
            title_cell.value     = f"{course.code} — {course.name}"
            title_cell.font      = Font(name='Calibri', bold=True, color='FFFFFF', size=13)
            title_cell.fill      = meta_fill
            title_cell.alignment = center
            title_cell.border    = border
            ws.row_dimensions[1].height = 24

            ws.merge_cells('A2:G2')
            meta2 = ws['A2']
            meta2.value     = (
                f"Programme: {course.programme.name if course.programme else '—'}  |  "
                f"Threshold: {threshold}%  |  "
                f"Enrolled: {enrollments.count()}  |  "
                f"Sessions held: {AttendanceSession.objects.filter(course=course).count()}  |  "
                f"Generated: {now.strftime('%d %b %Y')}"
            )
            meta2.font      = Font(name='Calibri', size=9, color='4F46E5')
            meta2.alignment = center
            meta2.border    = border
            ws.row_dimensions[2].height = 16

            col_headers = ["#", "Student Name", "Student ID", "Email", "Sessions Attended", "Sessions Held", "Attendance %"]
            col_widths  = [4,    28,             16,           30,      18,                  14,              14]
            for i, w in enumerate(col_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w

            cur_row = 3

            # ── Above threshold ──
            ws.merge_cells(f'A{cur_row}:G{cur_row}')
            sec = ws[f'A{cur_row}']
            sec.value     = f"✓ ABOVE THRESHOLD ({len(above)} students)"
            sec.font      = hdr_font
            sec.fill      = above_fill
            sec.alignment = center
            sec.border    = border
            ws.row_dimensions[cur_row].height = 18
            cur_row += 1

            _hdr_row(ws, cur_row, col_headers, above_fill)
            ws.row_dimensions[cur_row].height = 18
            cur_row += 1

            for idx, enr in enumerate(above, 1):
                pct = float(enr.attendance_percentage)
                _data_row(ws, cur_row, [
                    idx, enr.student.get_full_name(),
                    enr.student.student_id or "—",
                    enr.student.email,
                    enr.attended_sessions, enr.total_sessions,
                    f"{pct:.1f}%",
                ], idx % 2 == 0)
                cur_row += 1

            cur_row += 1  # blank separator

            # ── Below threshold ──
            ws.merge_cells(f'A{cur_row}:G{cur_row}')
            sec2 = ws[f'A{cur_row}']
            sec2.value     = f"✗ BELOW THRESHOLD — AT RISK ({len(below)} students)"
            sec2.font      = hdr_font
            sec2.fill      = below_fill
            sec2.alignment = center
            sec2.border    = border
            ws.row_dimensions[cur_row].height = 18
            cur_row += 1

            _hdr_row(ws, cur_row, col_headers, below_fill)
            ws.row_dimensions[cur_row].height = 18
            cur_row += 1

            for idx, enr in enumerate(below, 1):
                pct = float(enr.attendance_percentage)
                _data_row(ws, cur_row, [
                    idx, enr.student.get_full_name(),
                    enr.student.student_id or "—",
                    enr.student.email,
                    enr.attended_sessions, enr.total_sessions,
                    f"{pct:.1f}%",
                ], idx % 2 == 0)
                # Colour attendance % cell red
                pct_cell = ws.cell(row=cur_row, column=7)
                pct_cell.font = Font(name='Calibri', size=10, bold=True, color='DC2626')
                cur_row += 1

        if not wb.sheetnames:
            ws = wb.create_sheet("No Data")
            ws['A1'] = "No courses with enrollments found."

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"threshold_report_{now.strftime('%Y%m%d_%H%M')}.xlsx"
        resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{fname}"'
        return resp

